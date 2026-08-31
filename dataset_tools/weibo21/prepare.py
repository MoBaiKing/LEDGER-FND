from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from dataset_tools.common import (
    SPLITS,
    require_columns,
    verify_image,
    write_workspace,
)
from mmfnd.utils import resolve_path


def _read_sheet(
    path: Path,
    image_root: Path,
    split: str,
    verify_decode: bool,
) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = ["" if value is None else str(value).strip() for value in next(rows)]
    require_columns(set(header), {"label", "content", "image"}, path)
    index = {name: header.index(name) for name in header if name}

    records = []
    for excel_row, row in enumerate(rows, start=2):
        label = int(row[index["label"]])
        raw_images = str(row[index["image"]] or "")
        image_paths = [
            item.strip().replace("\\", "/")
            for item in raw_images.split("|")
            if item.strip()
        ]
        if not image_paths:
            raise ValueError(f"{path.name}:{excel_row} 图片字段为空")
        hashes = [
            verify_image(image_root / relative, verify_decode)
            for relative in image_paths
        ]
        expected_dir = "rumor_images" if label == 0 else "nonrumor_images"
        records.append({
            "id": f"{split}-{excel_row - 2:06d}",
            "excel_row": excel_row,
            "label": label,
            "text": str(row[index["content"]] or "").strip(),
            "images": image_paths,
            "image_sha256": hashes,
            "category": (
                str(row[index["category"]] or "").strip()
                if "category" in index else ""
            ),
            "title": (
                str(row[index["title"]] or "").strip()
                if "title" in index else ""
            ),
            "source": (
                str(row[index["source"]] or "").strip()
                if "source" in index else ""
            ),
            "folder_label_consistent": all(
                Path(item).parts[0] == expected_dir for item in image_paths
            ),
        })
    workbook.close()
    return records


def prepare(project_root: Path, config: dict, verify_decode: bool) -> Path:
    annotations_root = resolve_path(project_root, config["data"]["root"])
    image_root = resolve_path(project_root, config["data"]["image_root"])
    records_by_split: dict[str, list[dict]] = {}
    hash_uses: dict[str, list[dict]] = defaultdict(list)

    for split in SPLITS:
        records = _read_sheet(
            annotations_root / f"{split}_datasets.xlsx",
            image_root,
            split,
            verify_decode,
        )
        records_by_split[split] = records
        for record in records:
            for image_hash, image_path in zip(
                record["image_sha256"], record["images"]
            ):
                hash_uses[image_hash].append({
                    "split": split,
                    "id": record["id"],
                    "image": image_path,
                })

    overlaps = [
        {"sha256": value, "uses": uses}
        for value, uses in hash_uses.items()
        if len({use["split"] for use in uses}) > 1
    ]
    report = {
        "adapter": "dataset_tools.weibo21.prepare",
        "annotations_root": str(annotations_root.resolve()),
        "image_root": str(image_root.resolve()),
        "cross_split_duplicate_images": len(overlaps),
        "overlap_details": overlaps,
        "note": (
            "官方划分保持不变；跨分区重复图片只报告、不静默删除。"
        ),
    }
    return write_workspace(project_root, config, records_by_split, report)
