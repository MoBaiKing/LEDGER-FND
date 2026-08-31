from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
import json
from pathlib import Path

from openpyxl import load_workbook
from sklearn.model_selection import train_test_split

from dataset_tools.common import (
    ImageVerifier, dump_json, normalize_text, require_columns, resolve_path,
    safe_relative_path, sha256, text_sha256, write_workspace,
)


def _split(records: list[dict], seed: int) -> dict[str, list[dict]]:
    groups = defaultdict(list)
    for record in records:
        groups[(record["text_sha256"], record["image_sha256"][0])].append(record)
    group_keys = sorted(groups)
    group_labels = [
        Counter(int(row["label"]) for row in groups[key]).most_common(1)[0][0]
        for key in group_keys
    ]
    train_keys, temporary_keys = train_test_split(
        group_keys, test_size=0.2, random_state=seed, stratify=group_labels,
    )
    temporary_labels = [
        Counter(int(row["label"]) for row in groups[key]).most_common(1)[0][0]
        for key in temporary_keys
    ]
    val_keys, test_keys = train_test_split(
        temporary_keys, test_size=0.5, random_state=seed,
        stratify=temporary_labels,
    )
    expand = lambda keys: sorted(
        (row for key in keys for row in groups[key]),
        key=lambda row: int(row["source_row"]),
    )
    return {
        "train": expand(train_keys),
        "val": expand(val_keys),
        "test": expand(test_keys),
    }


def prepare(project_root: Path, config: dict, verify_decode: bool) -> Path:
    data_root = resolve_path(project_root, config["data"]["root"])
    image_root = resolve_path(project_root, config["data"]["image_root"])
    source = data_root / config["data"].get("source_file", "gossipcop_all_data.xlsx")
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = ["" if value is None else str(value).strip() for value in next(rows)]
    require_columns(set(header), {"id", "image_name", "l1bel", "text"}, source)
    index = {name: header.index(name) for name in header if name}
    verifier, records, exclusions, seen_ids = ImageVerifier(verify_decode), [], [], set()
    source_rows = 0
    for source_row, row in enumerate(rows, start=2):
        source_rows += 1
        source_id = normalize_text(row[index["id"]])
        raw_label = row[index["l1bel"]]
        text = normalize_text(row[index["text"]])
        image_name = normalize_text(row[index["image_name"]])
        reason = None
        if not source_id:
            reason = "empty_id"
        elif source_id in seen_ids:
            reason = "duplicate_source_id"
        elif raw_label not in (0, 1):
            reason = "non_binary_label"
        elif not text:
            reason = "empty_text"
        elif not image_name:
            reason = "empty_image"
        if reason:
            exclusions.append({"source_id": source_id, "source_row": source_row, "reason": reason})
            continue
        seen_ids.add(source_id)
        try:
            relative = safe_relative_path(f"gossipcop_images/{Path(image_name).name}")
            image_hash = verifier(image_root / relative)
        except Exception as error:
            exclusions.append({"source_id": source_id, "source_row": source_row, "reason": f"invalid_image:{type(error).__name__}"})
            continue
        label = 1 - int(raw_label)
        records.append({
            "id": f"gossipcop-{source_id}", "source_id": source_id,
            "source_row": source_row, "raw_label": int(raw_label),
            "label": label, "label_name": "fake" if label == 0 else "real",
            "text": text, "text_sha256": text_sha256(text),
            "images": [relative], "image_sha256": [image_hash],
            "category": "entertainment",
            "cleaning_version": "gossipcop_preserve_samples_grouped_8_1_1_v3",
        })
    workbook.close()
    seed = int(config.get("seed", 200408))
    records_by_split = _split(records, seed)
    split_sets = {
        split: {
            "id": {row["id"] for row in values},
            "sample": {
                (row["text_sha256"], row["image_sha256"][0])
                for row in values
            },
        } for split, values in records_by_split.items()
    }
    leakage = {
        f"{left}_{right}_shared_{key}": len(split_sets[left][key] & split_sets[right][key])
        for left, right in (("train", "val"), ("train", "test"), ("val", "test"))
        for key in ("id", "sample")
    }
    if any(leakage.values()):
        raise RuntimeError(f"GossipCop 清洗后仍有跨分区泄漏: {leakage}")
    report = {
        "adapter": "dataset_tools.gossipcop.prepare",
        "protocol": "gossipcop_preserve_samples_grouped_8_1_1_v3",
        "source_file": str(source.resolve()), "source_sha256": sha256(source),
        "image_root": str(image_root.resolve()),
        "label_semantics": {"source_0": "real", "source_1": "fake", "0": "fake", "1": "real"},
        "label_policy": "preserve every source-row l1bel; canonical label = 1 - l1bel",
        "sample_policy": "preserve every valid source row; exact text+image copies stay in one split",
        "seed": seed, "source_rows": source_rows, "accepted_samples": len(records),
        "excluded_samples": len(exclusions),
        "excluded_reasons": dict(Counter(item["reason"] for item in exclusions)),
        "text_cleaning": "HTML entity decode + Unicode NFKC + zero-width removal + whitespace normalization",
        "verify_decode": verify_decode, "validated_unique_images": verifier.unique_images,
        "leakage_checks": leakage,
        "completed_at": datetime.now().astimezone().isoformat(),
    }
    output = write_workspace(project_root, config, records_by_split, report)
    (output / "dataset_audit.json").unlink(missing_ok=True)
    with (output / "cleaning_exclusions.jsonl").open("w", encoding="utf-8") as file:
        for exclusion in exclusions:
            file.write(json.dumps(exclusion, ensure_ascii=False) + "\n")
    dump_json({"dataset": "gossipcop", "splits": {split: len(values) for split, values in records_by_split.items()}}, output / "cleaning_summary.json")
    return output
